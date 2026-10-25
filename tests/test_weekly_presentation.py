import os
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from unittest.mock import MagicMock, patch

from lxml import etree
from openpyxl import Workbook, load_workbook

from weekly.presentation import (
    ChartData,
    ChartOverflow,
    _apply_presentation_review,
    _month_key,
    _metric_value,
    _rounded,
    _update_embedded,
    _update_chart_xml,
    _chart_axis_overflows,
    _set_chart_overflow_note,
    extract_mobile_charts,
    write_weekly_presentation,
)
from weekly.presentation_review import load_presentation_review


class WeeklyPresentationTests(unittest.TestCase):
    def test_chart_metrics_keep_precision_and_reject_nonfinite_values(self) -> None:
        self.assertEqual(_metric_value(2.5714285714285716), 2.5714285714285716)
        self.assertEqual(_metric_value(100.125), 100.125)
        self.assertIsInstance(_metric_value(100.0), int)
        for invalid in (True, None, "2.5", float("nan"), float("inf")):
            with self.subTest(value=invalid), self.assertRaises(ValueError):
                _metric_value(invalid)

    def test_extract_uses_unrounded_pivot_values_for_chart_data(self) -> None:
        def monthly(*values):
            return {"Sep-26": values}, values
        def weekly(value):
            return {"W39": (value,)}, (value,)
        workbook = MagicMock()
        workbook.__contains__.return_value = True
        with patch("weekly.presentation.load_workbook", return_value=workbook), \
             patch("weekly.presentation._section", return_value=1), \
             patch("weekly.presentation._pivot_rows", side_effect=[
                 monthly(1234.567), monthly(12), monthly(1.25, 2.375),
                 weekly(100.125), weekly(0.524), weekly(7), weekly(5),
                 weekly(0.25), weekly(2.5714285714285716), weekly(1.75), weekly(1.125),
             ]):
            charts, averages = extract_mobile_charts("mobile.xlsx", date(2026, 9, 30), 39)
        self.assertEqual(charts[1].first, (0.25,))
        self.assertEqual(charts[1].second, (1.75,))
        self.assertEqual(charts[2].second, (100.125,))
        self.assertEqual(charts[3].second, (0.524,))
        self.assertEqual(charts[4].first, (2.5714285714285716,))
        self.assertEqual(charts[5].second, (2.375,))
        self.assertEqual(charts[6].second, (1234.567,))
        self.assertEqual(averages["bulk_cash"], 2)

    def test_fixed_axis_limits_are_preserved_and_overflow_uses_the_correct_axis(self) -> None:
        ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        root = etree.fromstring(f'''<c:chartSpace xmlns:c="{ns}"><c:plotArea>
          <c:lineChart><c:ser><c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>134</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser><c:axId val="1"/></c:lineChart>
          <c:lineChart><c:ser><c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>3</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser><c:axId val="2"/></c:lineChart>
          <c:valAx><c:axId val="1"/><c:scaling><c:min val="0"/><c:max val="10"/></c:scaling></c:valAx>
          <c:valAx><c:axId val="2"/><c:scaling><c:min val="0"/><c:max val="10"/></c:scaling></c:valAx>
        </c:plotArea></c:chartSpace>''')
        original = etree.tostring(root)
        outliers = _chart_axis_overflows(root)
        namespace = {"c": ns}
        self.assertEqual(root.xpath("//c:valAx[c:axId/@val='1']/c:scaling/c:max/@val", namespaces=namespace), ["10"])
        self.assertEqual(root.xpath("//c:valAx[c:axId/@val='2']/c:scaling/c:max/@val", namespaces=namespace), ["10"])
        self.assertEqual(root.xpath("//c:valAx/c:scaling/c:min/@val", namespaces=namespace), ["0", "0"])
        self.assertEqual(outliers, [ChartOverflow("", "Point 1", 134.0, 10.0, "max")])
        self.assertEqual(etree.tostring(root), original)

    def test_overflow_note_is_editable_and_does_not_change_plot_layout_or_values(self) -> None:
        ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
        root = etree.fromstring(f'''<c:chartSpace xmlns:c="{ns}" xmlns:a="{drawing}">
          <c:chart><c:autoTitleDeleted val="1"/><c:plotArea><c:layout/>
          <c:catAx><c:lblOffset val="100"/></c:catAx></c:plotArea></c:chart></c:chartSpace>''')
        plot = etree.tostring(root.find(f".//{{{ns}}}plotArea"))
        _set_chart_overflow_note(root, [
            ChartOverflow("LEADTIME CASH-COMPLETE (SAD)", "W40", 134.0, 11.0, "max")
        ], weekly=True)
        namespace = {"c": ns, "a": drawing}
        self.assertEqual(root.xpath("//c:title//a:t/text()", namespaces=namespace),
                         ["CASH-COMPLETE W40", "134 (>11, off-scale)"])
        self.assertEqual(root.xpath("//c:title/c:overlay/@val", namespaces=namespace), ["1"])
        self.assertEqual(root.xpath("//c:catAx/c:lblOffset/@val", namespaces=namespace), ["100"])
        self.assertEqual(etree.tostring(root.find(f".//{{{ns}}}plotArea")), plot)
        _set_chart_overflow_note(root, [
            ChartOverflow("LEADTIME FAB-CASH (BRA)", "W41", 20.0, 10.0, "max")
        ], weekly=True)
        self.assertEqual(len(root.xpath("//c:title", namespaces=namespace)), 1)
        self.assertNotIn("W40", etree.tostring(root).decode())
        _set_chart_overflow_note(root, [], weekly=True)
        self.assertFalse(root.xpath("//c:title", namespaces=namespace))
        self.assertEqual(root.xpath("//c:autoTitleDeleted/@val", namespaces=namespace), ["1"])

    def test_monthly_overflow_summary_and_lower_bound_do_not_hide_other_outliers(self) -> None:
        ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
        root = etree.fromstring(f'<c:chartSpace xmlns:c="{ns}"><c:chart/></c:chartSpace>')
        points = [ChartOverflow("BULK SERVICE", "OCT-26", 134.0, 10.0, "max"),
                  ChartOverflow("MOBILE SERVICE", "SEP-26", -4.0, 0.0, "min"),
                  ChartOverflow("BULK SERVICE", "AUG-26", 12.0, 10.0, "max")]
        _set_chart_overflow_note(root, points, weekly=False)
        text = " ".join(root.xpath("//a:t/text()", namespaces={"a": drawing}))
        self.assertIn("OCT-26 BULK SERVICE = 134 (>10)", text)
        self.assertIn("SEP-26 MOBILE SERVICE = -4 (<0)", text)
        self.assertIn("+1 point", text)

    def test_existing_non_generated_title_is_preserved(self) -> None:
        ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        root = etree.fromstring(f'<c:chartSpace xmlns:c="{ns}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart><c:title><c:tx><a:t>Off-scale owner commentary</a:t></c:tx></c:title></c:chart></c:chartSpace>')
        original = etree.tostring(root)
        _set_chart_overflow_note(root, [], weekly=True)
        self.assertEqual(etree.tostring(root), original)
        with self.assertRaisesRegex(ValueError, "existing title"):
            _set_chart_overflow_note(root, [ChartOverflow("Qty", "W40", 100.0, 10.0, "max")], weekly=True)

    def test_attributed_chart_review_updates_a_point(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "presentation-review.csv"
            path.write_text(
                "chart,series,label,value,reviewed_by,review_reason\n"
                "4,2,W37,2,Dewi,Approved historical correction\n",
                encoding="utf-8",
            )
            decisions = load_presentation_review(path)
            data = ChartData(("W36", "W37", "W38"), (1, 2, 3), (1, 1, 2))
            self.assertEqual(
                _apply_presentation_review(4, data, decisions).second, (1, 2, 2)
            )
            with self.assertRaisesRegex(ValueError, "absent"):
                _apply_presentation_review(
                    4, ChartData(("W38",), (3,), (2,)), decisions
                )

    def test_chart_cache_refreshes_historical_and_current_points_from_pivot(self) -> None:
        ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        series = "".join(
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt idx=\"0\"><c:v>{title}</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:f>Sheet1!B1:D1</c:f><c:strCache><c:ptCount val=\"3\"/>{''.join(f'<c:pt idx=\"{i}\"><c:v>{label}</c:v></c:pt>' for i, label in enumerate(('W36', 'W37', 'W38')))}</c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:f>Sheet1!B2:D2</c:f><c:numCache><c:ptCount val=\"3\"/>{''.join(f'<c:pt idx=\"{i}\"><c:v>{value}</c:v></c:pt>' for i, value in enumerate(points))}</c:numCache></c:numRef></c:val></c:ser>"
            for title, points in (("LEADTIME FAB-CASH", (2, 3, 2)), ("LEADTIME CASH-COMPLETE", (2, 2, 2)))
        )
        xml = f'<c:chartSpace xmlns:c="{ns}">{series}</c:chartSpace>'.encode()
        current = ChartData(("W36", "W37", "W38", "W39"), (2, 2.5714285714285716, 2, 4.01), (2, 1.5, 2, 1.6))
        result = etree.fromstring(_update_chart_xml(xml, current, 3, 4))
        points = result.xpath("//c:ser/c:val//c:pt/c:v/text()", namespaces={"c": ns})
        self.assertEqual(points, [str(value) for value in (*current.first, *current.second)])
        self.assertEqual(result.xpath("//c:numCache/c:formatCode/text()", namespaces={"c": ns}),
                         ["#,##0", "#,##0"])

    def test_rounding_matches_excel_display(self) -> None:
        self.assertEqual(_rounded(2.5), 3)
        self.assertEqual(_rounded(1.49), 1)
        self.assertEqual(_month_key("Sep-26"), (2026, 9))
        self.assertIsNone(_month_key("Grand Total"))

    def test_embedded_weekly_data_keeps_history_and_appends_report_week(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        for column, label in enumerate(("W34", "W35", "W36", "W37"), start=2):
            sheet.cell(1, column, label)
        source = BytesIO()
        workbook.save(source)
        result, first = _update_embedded(
            source.getvalue(),
            ChartData(("W35", "W36", "W37", "W38"), (1, 2, 3, 4), (5, 6, 7, 8)),
            weekly=True,
            report_week=38,
        )
        self.assertEqual(first, 3)
        updated = load_workbook(BytesIO(result), read_only=True, data_only=True).active
        self.assertEqual(updated["B1"].value, "W34")
        self.assertEqual(updated["F1"].value, "W38")
        self.assertEqual(updated["F3"].value, 8)
        for row in (2, 3):
            self.assertEqual(updated.cell(row, 6).number_format, "#,##0")

        editable = load_workbook(BytesIO(result)).active
        self.assertTrue(editable.column_dimensions["B"].hidden)
        self.assertFalse(editable.column_dimensions["C"].hidden)
        next_result, next_first = _update_embedded(
            result,
            ChartData(("W36", "W37", "W38", "W39"), (2, 3, 4, 5), (6, 7, 8, 9)),
            weekly=True,
            report_week=39,
        )
        self.assertEqual(next_first, 4)
        next_sheet = load_workbook(BytesIO(next_result)).active
        self.assertTrue(next_sheet.column_dimensions["B"].hidden)
        self.assertTrue(next_sheet.column_dimensions["C"].hidden)
        self.assertEqual([next_sheet.cell(1, column).value for column in range(4, 8)],
                         ["W36", "W37", "W38", "W39"])
        self.assertFalse(next_sheet.column_dimensions["D"].hidden)

    def test_embedded_data_preserves_fractional_metrics_with_integer_display_format(self) -> None:
        source = BytesIO()
        workbook = Workbook()
        workbook.active.cell(1, 2, "W39")
        workbook.save(source)
        data = ChartData(("W39",), (2.5714285714285716,), (100.125,))
        result, first = _update_embedded(source.getvalue(), data, weekly=True, report_week=39)
        updated = load_workbook(BytesIO(result)).active
        self.assertAlmostEqual(updated.cell(2, first).value, data.first[0], places=12)
        self.assertEqual(updated.cell(3, first).value, 100.125)
        self.assertEqual(updated.cell(2, first).number_format, "#,##0")

    @unittest.skipUnless(os.environ.get("WEEKLY_REFERENCE_DIR"), "WEEKLY_REFERENCE_DIR not set")
    def test_w38_deck_uses_mobile_pivot_and_preserves_power_bi_slide(self) -> None:
        reference = Path(os.environ["WEEKLY_REFERENCE_DIR"])
        mobile = reference / "Master Raw Mobile Service - W38 2026.xlsx"
        template = reference / "Mobile Services W37 - 2026 (Weekly).pptx"
        charts, averages = extract_mobile_charts(mobile, date(2026, 9, 24), 38)
        self.assertEqual(charts[2].first, (39, 11, 21, 14))
        self.assertEqual(tuple(_rounded(value) for value in charts[4].first), (1, 2, 2, 2))
        self.assertEqual(charts[5].labels[-1], "SEP-26")
        self.assertEqual(averages["bulk_cash"], 3)
        with TemporaryDirectory() as directory:
            output = Path(directory) / "weekly.pptx"
            write_weekly_presentation(mobile_workbook=mobile, template_pptx=template,
                                      output_pptx=output, report_date=date(2026, 9, 24), report_week=38)
            with ZipFile(template) as before, ZipFile(output) as after:
                self.assertEqual(after.testzip(), None)
                self.assertEqual(before.read("ppt/media/image1.png"), after.read("ppt/media/image1.png"))
                slide = etree.fromstring(after.read("ppt/slides/slide2.xml"))
                drawing = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                self.assertIn("9/24/2026", slide.xpath("//a:t/text()", namespaces=drawing))
                self.assertNotIn("datetime1", after.read("ppt/slides/slide2.xml").decode())
                chart = etree.fromstring(after.read("ppt/charts/chart4.xml"))
                ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                self.assertEqual(chart.xpath("(//c:ser)[1]/c:val//c:pt/c:v/text()", namespaces=ns),
                                 [str(value) for value in charts[4].first])
                for index in range(4):
                    suffix = str(index) if index else ""
                    part = f"ppt/embeddings/Microsoft_Excel_Worksheet{suffix}.xlsx"
                    sheet = load_workbook(BytesIO(after.read(part))).active
                    visible_weeks = [sheet.cell(1, column).value
                                     for column in range(2, sheet.max_column + 1)
                                     if isinstance(sheet.cell(1, column).value, str)
                                     and sheet.cell(1, column).value.startswith("W")
                                     and not any(dimension.hidden and dimension.min <= column <= dimension.max
                                                 for dimension in sheet.column_dimensions.values()
                                                 if dimension.min is not None and dimension.max is not None)]
                    self.assertEqual(visible_weeks, ["W35", "W36", "W37", "W38"])
                    next_data, _ = _update_embedded(
                        after.read(part),
                        ChartData(("W36", "W37", "W38", "W39"), (1, 2, 3, 4), (5, 6, 7, 8)),
                        weekly=True,
                        report_week=39,
                    )
                    next_sheet = load_workbook(BytesIO(next_data)).active
                    next_visible = [next_sheet.cell(1, column).value
                                    for column in range(2, next_sheet.max_column + 1)
                                    if isinstance(next_sheet.cell(1, column).value, str)
                                    and next_sheet.cell(1, column).value.startswith("W")
                                    and not any(dimension.hidden and dimension.min <= column <= dimension.max
                                                for dimension in next_sheet.column_dimensions.values()
                                                if dimension.min is not None and dimension.max is not None)]
                    self.assertEqual(next_visible, ["W36", "W37", "W38", "W39"])


if __name__ == "__main__":
    unittest.main()
