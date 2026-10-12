import unittest
from datetime import date, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter

from generate_ide_tracking import (
    ACTUAL_RFS_DATE_HEADER,
    DEPT_HEADER,
    NEW_RFS_INITIAL_HEADER,
    QUOTE_ID_HEADER,
    build_all_order,
    apply_current_date_first_match,
    apply_date_formats,
    apply_previous_field_fallbacks,
    apply_text_formats,
    completion_thresholds,
    date_from_filename,
    financial_value_or_zero,
    has_expected_all_order_freeze_panes,
    hidden_row_ranges_between_sections,
    identify_target_complete_pivots,
    is_missing,
    load_collabs_lookup,
    load_source_new_rfs_lookup,
    normalize_month_first_dashboard_dates,
    normalize_pivot_section_row_visibility,
    parse_excel_datetime,
    reconcile_date_column,
    reconciled_date_value,
    resolve_financial_values,
    target_period,
    target_value,
    target_value_with_previous,
    validate_output,
    validate_raw_headers,
)


class IdeDynamicColumnsTests(unittest.TestCase):
    RAW_HEADERS = (
        "NO", "QUOTE ID", "PROJECT NAME", "QUOTE NAME", "QUOTE NAME 2",
        "CUSTOMER NAME", "PROJECT TYPE", "BRAND", "PARTNERS",
        "STATUS Catalog / non Catalog", "NOMOR PO", "SO", "CID", "QUOTE DATE",
        "FAB SIGN DATE", "FAB UPLOAD DATE", "NEW RFS INITIAL", "CASH DATE",
        "SO DATE", "SERVICE READY", "ACTUAL RFS DATE", "AOS SENT DATE",
        "AOS SIGN DATE", "SO COMPLETION DATE", "STATUS DELIVERY", "SD Note",
        "SALES", "SEGMENT SALES", "PRODUCT2", "ADMIN", "APM", "PM", "DEPT.",
        "PIC CUSTOMER", "CUSTOMER ADDRESS", "Project Category", "REVENUE",
        "REVENUE UPSIDE", "PENALTY BC", "CoS BC", "RISK BC", "PENALTY ACTUAL",
        "COS ACTUAL (PO)", "RISK ACTUAL", "COST PLANNED", "ACTUAL COST",
        "COST EFFICIENCY INDEX (CEI)", "STATUS HOTO D2", "OTD",
    )

    def setUp(self) -> None:
        raw_row = dict.fromkeys(self.RAW_HEADERS)
        raw_row.update({
            "NO": 1, "QUOTE ID": "Q-1", "QUOTE NAME": "Example",
            "FAB UPLOAD DATE": datetime(2026, 10, 1),
            "NEW RFS INITIAL": datetime(2026, 10, 12),
            "ACTUAL RFS DATE": datetime(2026, 10, 12),
            "STATUS DELIVERY": "SO Complete",
        })
        self.raw = pd.DataFrame([raw_row])
        self.previous = pd.DataFrame([{
            "QUOTE ID": "Q-1", "ORDER TYPE": "New Registration",
            "OTC": 100, "MRC": 200, "STATUS DELIVERY": "SO Complete",
            "Pre-Installation Start Date": datetime(2026, 10, 1),
        }])

    def _build(self, raw):
        with patch("generate_ide_tracking.load_collabs_lookup", return_value={}):
            return build_all_order(
                raw, self.previous, Path("collabs.csv"), date(2026, 10, 12),
                Path("Daily Tracking IDE 9 October 2026.xlsx"),
            ).dataframe

    def test_raw_additions_keep_their_values_and_existing_business_columns(self) -> None:
        baseline = self._build(self.raw)
        self.assertEqual(len(baseline.columns), 61)
        for position in (6, len(self.raw.columns)):
            with self.subTest(position=position):
                extended = self.raw.copy()
                extended.insert(position, "MONTH COMPLETE", ["Okt"])
                extended["SOURCE NOTES"] = ["Confirmed"]
                result = self._build(extended)
                self.assertEqual(len(result.columns), 63)
                self.assertEqual(result["MONTH COMPLETE"].tolist(), ["Okt"])
                self.assertEqual(result["SOURCE NOTES"].tolist(), ["Confirmed"])
                pd.testing.assert_frame_equal(
                    result.drop(columns=["MONTH COMPLETE", "SOURCE NOTES"]), baseline
                )

    def test_missing_required_column_has_a_specific_error(self) -> None:
        with self.assertRaisesRegex(ValueError, "missing required columns: PARTNERS"):
            self._build(self.raw.drop(columns=["PARTNERS"]))

    def test_duplicate_normalized_headers_are_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "duplicate columns: QUOTE ID"):
            validate_raw_headers([*self.RAW_HEADERS, "  QUOTE   ID  "])

    def test_extra_column_preserves_date_like_text_and_formats_typed_dates(self) -> None:
        ranges = {}

        class Worksheet:
            def Range(self, address):
                return ranges.setdefault(address, SimpleNamespace(NumberFormat="General"))

        dataframe = pd.DataFrame({
            "MONTH COMPLETE": ["4/30/2026", datetime(2026, 4, 30), None, 7],
            "SOURCE NOTES": ["A", "B", None, "3/4/2026"],
            "COUNT": [0, 1, None, 2],
        }, dtype=object)
        worksheet = Worksheet()
        apply_text_formats(worksheet, dataframe)
        apply_date_formats(worksheet, dataframe)
        self.assertEqual(ranges["A2:A2"].NumberFormat, "@")
        self.assertEqual(ranges["A3:A3"].NumberFormat, "mm/dd/yy")
        self.assertEqual(ranges["B2:B3"].NumberFormat, "@")
        self.assertEqual(ranges["B5:B5"].NumberFormat, "@")
        self.assertFalse(any(address.startswith("C") for address in ranges))

    def _output_workbook(self, directory):
        headers = ["QUOTE ID", "OTC", "MRC", "YEAR FAB UPLOAD"]
        headers.extend(f"Source field {index}" for index in range(1, 58))
        headers.append("MONTH COMPLETE")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ALL ORDER"
        sheet.append(headers)
        sheet.append(["Q-1", 100, 200, "2026", *([None] * 57), "Okt"])
        sheet.cell(2, 4).number_format = "@"
        sheet.freeze_panes = "C1"
        sheet.add_table(Table(displayName="Table47", ref=f"A1:{get_column_letter(len(headers))}2"))
        workbook.create_sheet("PIVOT")
        return workbook, Path(directory) / "output.xlsx", headers

    def test_output_accepts_source_column_growth_and_ignores_unused_cell_formatting(self) -> None:
        with TemporaryDirectory() as directory:
            workbook, path, headers = self._output_workbook(directory)
            workbook["ALL ORDER"].cell(2, 70).number_format = "#,##0"
            workbook.save(path)
            workbook.close()
            result = validate_output(path, ["Q-1"], 0, expected_headers=headers)
            self.assertEqual(result["columns"], 62)
            self.assertEqual(result["table_ref"], "A1:BJ2")

    def test_output_rejects_changed_header_even_when_column_count_matches(self) -> None:
        with TemporaryDirectory() as directory:
            workbook, path, headers = self._output_workbook(directory)
            workbook["ALL ORDER"].cell(1, len(headers)).value = "OTHER FIELD"
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "column names or order"):
                validate_output(path, ["Q-1"], 0, expected_headers=headers)

    def test_output_table_must_include_the_added_source_column(self) -> None:
        with TemporaryDirectory() as directory:
            workbook, path, headers = self._output_workbook(directory)
            workbook["ALL ORDER"].tables["Table47"].ref = "A1:BI2"
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "Table47 range"):
                validate_output(path, ["Q-1"], 0, expected_headers=headers)

    def test_output_rejects_a_pivot_cache_that_did_not_pick_up_new_columns(self) -> None:
        with TemporaryDirectory() as directory:
            workbook, path, headers = self._output_workbook(directory)
            workbook["PIVOT"]._pivots = [SimpleNamespace(
                name="StalePivot",
                cache=SimpleNamespace(cacheFields=[
                    SimpleNamespace(name=header) for header in headers[:-1]
                ]),
            )]
            with patch("generate_ide_tracking.load_workbook", return_value=workbook):
                with self.assertRaisesRegex(ValueError, "StalePivot.*source columns"):
                    validate_output(path, ["Q-1"], 1, expected_headers=headers)


class FakeRangeDimension:
    def __init__(self, count):
        self.Count = count


class FakeTableRange:
    def __init__(self, row, column, rows, columns):
        self.Row = row
        self.Column = column
        self.Rows = FakeRangeDimension(rows)
        self.Columns = FakeRangeDimension(columns)


class FakePivot:
    def __init__(self, name, row, column, rows, columns, row_range=None):
        self.Name = name
        self.TableRange2 = FakeTableRange(row, column, rows, columns)
        start, count = row_range if row_range is not None else (row, rows)
        self.RowRange = FakeTableRange(start, column, count, 1)


class FakePivotCollection:
    def __init__(self, pivots):
        self._pivots = pivots
        self.Count = len(pivots)

    def __call__(self, index):
        return self._pivots[index - 1]


class FakePivotWorksheet:
    def __init__(self, pivots):
        self._pivots = FakePivotCollection(pivots)

    def PivotTables(self):
        return self._pivots


class FakeUsedRange:
    def __init__(self, values, row=1):
        self.Value2 = values
        self.Row = row


class FakeHiddenRows:
    def __init__(self, worksheet, start_row, end_row):
        self.worksheet = worksheet
        self.start_row = start_row
        self.end_row = end_row

    @property
    def EntireRow(self):
        return self

    @property
    def Hidden(self):
        return None

    @Hidden.setter
    def Hidden(self, value):
        self.worksheet.visibility_changes.append(
            (self.start_row, self.end_row, bool(value))
        )


class FakeVisibilityWorksheet(FakePivotWorksheet):
    def __init__(self, pivots, values):
        super().__init__(pivots)
        self.UsedRange = FakeUsedRange(values)
        self.visibility_changes = []

    def Cells(self, row, column):
        return row, column

    def Range(self, start_cell, end_cell):
        return FakeHiddenRows(self, start_cell[0], end_cell[0])

    def Rows(self, row):
        return FakeHiddenRows(self, row, row)


class IdeDateParsingTests(unittest.TestCase):
    def test_refreshed_pivot_rows_are_unhidden_before_gap_rows(self) -> None:
        values = tuple(
            (value,)
            for value in [
                "TARGET COMPLETE",
                None,
                "header",
                "row 1",
                "row 2",
                "row 3",
                "row 4",
                "Grand Total",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "ORDER AGING START FROM PRE-INSTALLATION STATUS",
            ]
        )
        worksheet = FakeVisibilityWorksheet(
            [FakePivot("target-complete", 3, 1, 6, 4)],
            values,
        )

        normalize_pivot_section_row_visibility(worksheet)

        self.assertIn((3, 8, False), worksheet.visibility_changes)
        self.assertIn((1, 1, False), worksheet.visibility_changes)
        self.assertIn((20, 20, False), worksheet.visibility_changes)
        self.assertIn((10, 17, True), worksheet.visibility_changes)

    def test_pivot_filter_rows_remain_hidden(self) -> None:
        values = tuple(
            (value,)
            for value in [
                "TARGET COMPLETE",
                None,
                "New Registration",
                None,
                "TARGET  Detemined as 1 Sep 26",
                "YEAR FAB UPLOAD",
                "YEAR FAB UPLOAD",
                "Count of QUOTE ID",
                "Div./Dept.",
                "Order",
                "Grand Total",
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                "ORDER AGING START FROM PRE-INSTALLATION STATUS",
            ]
        )
        worksheet = FakeVisibilityWorksheet(
            [FakePivot("target-complete", 5, 4, 7, 7, row_range=(9, 3))],
            values,
        )

        normalize_pivot_section_row_visibility(worksheet)

        self.assertIn((9, 11, False), worksheet.visibility_changes)
        self.assertFalse(
            any(
                start <= row <= end and not hidden
                for start, end, hidden in worksheet.visibility_changes
                for row in range(5, 9)
            )
        )

    def test_freeze_pane_validation_ignores_horizontal_scroll_position(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.freeze_panes = "C1"
        worksheet.sheet_view.pane.topLeftCell = "N1"
        worksheet.sheet_view.pane.state = "frozenSplit"

        self.assertTrue(has_expected_all_order_freeze_panes(worksheet))

        worksheet.sheet_view.pane.xSplit = 2.5
        self.assertFalse(has_expected_all_order_freeze_panes(worksheet))

        worksheet.freeze_panes = "N1"
        self.assertFalse(has_expected_all_order_freeze_panes(worksheet))

    def test_hidden_row_ranges_follow_dynamic_section_boundaries(self) -> None:
        self.assertEqual(
            hidden_row_ranges_between_sections(25, 113),
            [(27, 110)],
        )
        self.assertEqual(
            hidden_row_ranges_between_sections(137, 166, {144}),
            [(139, 142), (145, 163)],
        )

    def test_target_complete_pivots_allow_dynamic_start_rows(self) -> None:
        summary = FakePivot("summary", 7, 1, 45, 3)
        status = FakePivot("status", 6, 4, 46, 7)
        later_left = FakePivot("later-left", 68, 1, 20, 5)
        later_right = FakePivot("later-right", 65, 8, 18, 4)
        worksheet = FakePivotWorksheet(
            [later_right, status, later_left, summary]
        )

        actual_summary, actual_status = identify_target_complete_pivots(worksheet)

        self.assertIs(actual_summary, summary)
        self.assertIs(actual_status, status)

    def test_target_complete_pivots_reject_unrelated_later_section(self) -> None:
        summary = FakePivot("summary", 7, 1, 45, 3)
        unrelated_status = FakePivot("later-right", 65, 8, 18, 4)
        worksheet = FakePivotWorksheet([summary, unrelated_status])

        with self.assertRaisesRegex(RuntimeError, "TARGET COMPLETE"):
            identify_target_complete_pivots(worksheet)

    def test_report_filename_accepts_english_month_with_any_case(self) -> None:
        expected = date(2026, 8, 18)

        self.assertEqual(date_from_filename(Path("IDE DASHBOARD 18 August 2026.xlsx")), expected)
        self.assertEqual(date_from_filename(Path("IDE DASHBOARD 18 august 2026.xlsx")), expected)
        self.assertEqual(date_from_filename(Path("IDE DASHBOARD 18 AUGUST 2026.xlsx")), expected)

    def test_report_filename_accepts_indonesian_month_with_any_case(self) -> None:
        expected = date(2026, 8, 18)

        self.assertEqual(date_from_filename(Path("IDE DASHBOARD 18 Agustus 2026.xlsx")), expected)
        self.assertEqual(date_from_filename(Path("IDE DASHBOARD 18 agustus 2026.xlsx")), expected)
        self.assertEqual(date_from_filename(Path("IDE DASHBOARD 18 AGUSTUS 2026.xlsx")), expected)

    def test_report_filename_accepts_english_and_indonesian_month_abbreviations(self) -> None:
        cases = {
            "IDE DASHBOARD 18 Aug 2026.xlsx": date(2026, 8, 18),
            "IDE DASHBOARD 18 AGU 2026.xlsx": date(2026, 8, 18),
            "IDE DASHBOARD 18 Agt 2026.xlsx": date(2026, 8, 18),
            "IDE DASHBOARD 18 Sep 2026.xlsx": date(2026, 9, 18),
            "IDE DASHBOARD 18 Sept 2026.xlsx": date(2026, 9, 18),
            "IDE DASHBOARD 18 Oct 2026.xlsx": date(2026, 10, 18),
            "IDE DASHBOARD 18 Okt 2026.xlsx": date(2026, 10, 18),
            "IDE DASHBOARD 18 Dec 2026.xlsx": date(2026, 12, 18),
            "IDE DASHBOARD 18 Des 2026.xlsx": date(2026, 12, 18),
        }

        for filename, expected in cases.items():
            with self.subTest(filename=filename):
                self.assertEqual(date_from_filename(Path(filename)), expected)

    def test_slash_dates_follow_dashboard_day_first_format(self) -> None:
        self.assertEqual(
            parse_excel_datetime("10/01/2026"),
            datetime(2026, 1, 10),
        )
        self.assertEqual(
            parse_excel_datetime("04/05/2026"),
            datetime(2026, 5, 4),
        )

    def test_iso_dates_are_unchanged(self) -> None:
        self.assertEqual(
            parse_excel_datetime("2026-08-13 07:30:00"),
            datetime(2026, 8, 13, 7, 30),
        )

    def test_month_first_can_be_selected_for_region_sensitive_fields(self) -> None:
        self.assertEqual(
            parse_excel_datetime("08/06/2026", dayfirst=False),
            datetime(2026, 8, 6),
        )
        self.assertEqual(
            parse_excel_datetime("22/06/2026", dayfirst=False),
            datetime(2026, 6, 22),
        )

    def test_region_sensitive_dashboard_dates_are_normalized_month_first(self) -> None:
        dataframe = pd.DataFrame(
            {
                NEW_RFS_INITIAL_HEADER: ["08/06/2026"],
                ACTUAL_RFS_DATE_HEADER: ["11/05/2026"],
            }
        )

        normalize_month_first_dashboard_dates(dataframe)

        self.assertEqual(dataframe.at[0, NEW_RFS_INITIAL_HEADER], datetime(2026, 8, 6))
        self.assertEqual(dataframe.at[0, ACTUAL_RFS_DATE_HEADER], datetime(2026, 11, 5))

    def test_duplicate_first_match_does_not_fill_an_empty_current_row(self) -> None:
        dataframe = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["A", "A", "B", "B"],
                NEW_RFS_INITIAL_HEADER: [
                    datetime(2026, 4, 15),
                    datetime(2026, 4, 30),
                    datetime(2026, 6, 8),
                    pd.NA,
                ],
                ACTUAL_RFS_DATE_HEADER: [
                    datetime(2026, 4, 15),
                    datetime(2026, 4, 30),
                    datetime(2026, 7, 1),
                    pd.NA,
                ],
            }
        )

        apply_current_date_first_match(dataframe)

        self.assertEqual(dataframe.at[1, NEW_RFS_INITIAL_HEADER], datetime(2026, 4, 30))
        self.assertEqual(dataframe.at[1, ACTUAL_RFS_DATE_HEADER], datetime(2026, 4, 15))
        self.assertTrue(is_missing(dataframe.at[3, NEW_RFS_INITIAL_HEADER]))
        self.assertTrue(is_missing(dataframe.at[3, ACTUAL_RFS_DATE_HEADER]))

    def test_source_new_rfs_lookup_uses_first_match_and_us_date_parsing(self) -> None:
        class WorksheetStub:
            def iter_rows(self, **_kwargs):
                return iter(
                    [
                        ("A", *([None] * 33), "08-06-2026 20:19:55"),
                        ("A", *([None] * 33), "09-06-2026 20:19:55"),
                    ]
                )

        lookup = load_source_new_rfs_lookup(WorksheetStub(), {"A"})

        self.assertEqual(lookup["A"], datetime(2026, 8, 6))

    def test_missing_current_date_is_not_filled_from_previous(self) -> None:
        value, used_fallback = reconciled_date_value(
            pd.NA,
            datetime(2026, 7, 1),
            date(2026, 8, 14),
        )

        self.assertTrue(is_missing(value))
        self.assertFalse(used_fallback)

    def test_actual_rfs_reconciliation_uses_first_quote_lookup_key(self) -> None:
        current = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1", "Q-1"],
                ACTUAL_RFS_DATE_HEADER: [
                    datetime(2025, 12, 16),
                    datetime(2025, 12, 16),
                ],
            }
        )
        previous = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1"],
                ACTUAL_RFS_DATE_HEADER: [datetime(2026, 12, 16)],
            }
        )

        values, fallback_count = reconcile_date_column(
            current,
            previous,
            ACTUAL_RFS_DATE_HEADER,
            date(2026, 8, 14),
        )

        self.assertEqual(list(values), [datetime(2026, 12, 16)] * 2)
        self.assertEqual(fallback_count, 2)

    def test_month_opening_carries_previous_rfs_dates(self) -> None:
        current = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1"],
                NEW_RFS_INITIAL_HEADER: [datetime(2026, 8, 9)],
            }
        )
        previous = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1"],
                NEW_RFS_INITIAL_HEADER: [datetime(2026, 9, 8)],
            }
        )

        values, fallback_count = reconcile_date_column(
            current,
            previous,
            NEW_RFS_INITIAL_HEADER,
            date(2026, 9, 1),
        )

        self.assertEqual(list(values), [datetime(2026, 9, 8)])
        self.assertEqual(fallback_count, 1)

    def test_non_opening_day_keeps_current_rfs_date(self) -> None:
        current = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1"],
                NEW_RFS_INITIAL_HEADER: [datetime(2026, 8, 9)],
            }
        )
        previous = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1"],
                NEW_RFS_INITIAL_HEADER: [datetime(2026, 9, 8)],
            }
        )

        values, fallback_count = reconcile_date_column(
            current,
            previous,
            NEW_RFS_INITIAL_HEADER,
            date(2026, 9, 2),
        )

        self.assertEqual(list(values), [datetime(2026, 8, 9)])
        self.assertEqual(fallback_count, 0)

    def test_month_opening_matches_duplicate_actual_rfs_by_occurrence(self) -> None:
        current = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1", "Q-1"],
                ACTUAL_RFS_DATE_HEADER: [
                    datetime(2026, 4, 15),
                    datetime(2026, 4, 15),
                ],
            }
        )
        previous = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["Q-1", "Q-1"],
                ACTUAL_RFS_DATE_HEADER: [
                    datetime(2026, 4, 15),
                    datetime(2026, 4, 30),
                ],
            }
        )

        values, fallback_count = reconcile_date_column(
            current,
            previous,
            ACTUAL_RFS_DATE_HEADER,
            date(2026, 9, 1),
        )

        self.assertEqual(
            list(values),
            [datetime(2026, 4, 15), datetime(2026, 4, 30)],
        )
        self.assertEqual(fallback_count, 1)

    def test_previous_field_fallback_uses_first_quote_match_and_preserves_na_error(self) -> None:
        current = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["A", "A", "B"],
                DEPT_HEADER: [pd.NA, pd.NA, "#N/A"],
            }
        )
        previous = pd.DataFrame(
            {
                QUOTE_ID_HEADER: ["A", "A"],
                DEPT_HEADER: ["First Dept", "Second Dept"],
            }
        )

        fallback_count = apply_previous_field_fallbacks(current, previous, [DEPT_HEADER])

        self.assertEqual(fallback_count, 2)
        self.assertEqual(current[DEPT_HEADER].iloc[:2].tolist(), ["First Dept", "First Dept"])
        self.assertEqual(current.at[2, DEPT_HEADER], "#N/A")

    def test_invalid_date_is_not_coerced(self) -> None:
        self.assertIsNone(parse_excel_datetime("24/26/2026"))

    def test_target_bucket_uses_parsed_calendar_month(self) -> None:
        report_date = date(2026, 8, 13)
        self.assertEqual(target_value("10/01/2026", report_date), "Before August")
        self.assertEqual(target_value("13/08/2026", report_date), "Target August")
        self.assertEqual(target_value("01/09/2026", report_date), "After August")

    def test_first_calendar_day_closes_the_previous_target_month(self) -> None:
        self.assertEqual(target_period(date(2026, 9, 1)), date(2026, 8, 31))
        self.assertEqual(target_period(date(2026, 1, 1)), date(2025, 12, 31))

    def test_later_report_days_use_the_current_target_month(self) -> None:
        self.assertEqual(target_period(date(2026, 9, 2)), date(2026, 9, 2))

    def test_missing_current_target_carries_the_previous_bucket_in_same_period(self) -> None:
        self.assertEqual(
            target_value_with_previous(
                pd.NA,
                "Target August",
                date(2026, 8, 31),
            ),
            "Target August",
        )

    def test_current_target_date_takes_precedence_over_previous(self) -> None:
        self.assertEqual(
            target_value_with_previous(
                datetime(2026, 9, 15),
                "Target Not Yet Inputted",
                date(2026, 8, 31),
            ),
            "After August",
        )

    def test_existing_target_is_carried_within_the_same_period(self) -> None:
        self.assertEqual(
            target_value_with_previous(
                datetime(2026, 8, 9),
                "After August",
                date(2026, 8, 31),
            ),
            "After August",
        )

    def test_existing_target_is_reclassified_when_period_changes(self) -> None:
        self.assertEqual(
            target_value_with_previous(
                pd.NA,
                "Target August",
                date(2026, 9, 2),
            ),
            "Target Not Yet Inputted",
        )

    def test_missing_previous_financial_value_falls_back_to_zero(self) -> None:
        self.assertEqual(financial_value_or_zero("1,250,000"), 1_250_000)
        self.assertEqual(financial_value_or_zero("#N/A"), 0)
        self.assertEqual(financial_value_or_zero(None), 0)

    def test_collabs_lookup_reads_only_requested_quote_ids(self) -> None:
        with TemporaryDirectory() as temp_dir:
            collabs_path = Path(temp_dir) / "collabs.csv"
            pd.DataFrame(
                {
                    "quote_num": ["Q-1", "Q-2"],
                    "otc": ["1,250", "9,999"],
                    "mrc": ["250", "999"],
                    "unused": ["x", "y"],
                }
            ).to_csv(collabs_path, index=False)

            lookup = load_collabs_lookup(collabs_path, {"Q-1"})

        self.assertEqual(lookup, {"Q-1": (1_250, 250)})

    def test_financial_fallback_uses_previous_then_collabs_then_zero(self) -> None:
        self.assertEqual(
            resolve_financial_values(100, 200, (300, 400)),
            (100, 200, False, False),
        )
        self.assertEqual(
            resolve_financial_values("#N/A", None, (300, 400)),
            (300, 400, True, False),
        )
        self.assertEqual(
            resolve_financial_values(None, None, (None, "#N/A")),
            (0, 0, False, True),
        )

    def test_completion_thresholds_follow_the_five_business_week_buckets(self) -> None:
        expectations = {
            1: (1, 30, 20, 20),
            6: (1, 30, 20, 20),
            7: (2, 40, 30, 30),
            13: (2, 40, 30, 30),
            14: (3, 50, 40, 40),
            20: (3, 50, 40, 40),
            21: (4, 60, 50, 50),
            28: (4, 60, 50, 50),
            29: (5, 60, 50, 50),
            31: (5, 60, 50, 50),
        }

        for day, expected in expectations.items():
            with self.subTest(day=day):
                thresholds = completion_thresholds(date(2026, 8, day))
                self.assertEqual(
                    (
                        thresholds.week,
                        thresholds.green_percent,
                        thresholds.yellow_percent,
                        thresholds.red_percent,
                    ),
                    expected,
                )


if __name__ == "__main__":
    unittest.main()
