import unittest
from datetime import date
from io import BytesIO
from unittest.mock import MagicMock, patch

from lxml import etree
from openpyxl import Workbook, load_workbook

from weekly.constants import MASTER_COLUMNS
from weekly.m2m_presentation import EXPECTED_SERIES, extract_m2m_charts
from weekly.presentation import ChartData, _update_chart_xml, _update_embedded
from weekly.presentation import _set_annotation
from weekly.profiles import report_profile
from weekly.workbook import _write_table


class WeeklyM2mTests(unittest.TestCase):
    def test_invalid_report_week_is_rejected_before_reading_a_workbook(self):
        for week in (0,54):
            with self.subTest(week=week), self.assertRaisesRegex(ValueError,"Report week"):
                extract_m2m_charts("missing.xlsx",date(2026,10,1),week)

    def test_profiles_keep_legacy_defaults_and_m2m_specific_structure(self):
        self.assertEqual(report_profile().name, "master")
        self.assertEqual(report_profile(mobile=True).name, "mobile")
        profile = report_profile("m2m")
        self.assertEqual(profile.services, {"M2M"})
        self.assertEqual(profile.summary_label_column, 17)
        self.assertTrue(profile.weekly_window)
        self.assertTrue(profile.preserve_extra_columns)
        self.assertFalse(profile.removed_pivots)
        with self.assertRaises(ValueError):
            report_profile("m2m", mobile=True)
        with self.assertRaises(ValueError):
            report_profile("unknown")

    def test_m2m_blank_structural_columns_are_not_written_with_short_arrays(self):
        sheet = MagicMock()
        cells, ranges = {}, {}
        sheet.Cells.side_effect = lambda r,c: cells.setdefault((r,c), MagicMock())
        def region(first, last):
            start = next(key for key,value in cells.items() if value is first)
            end = next(key for key,value in cells.items() if value is last)
            area = ranges.setdefault((start,end), MagicMock())
            area.Value2 = ((None,None,None,None),)
            return area
        sheet.Range.side_effect = region
        table = sheet.ListObjects.return_value
        table.ListColumns.Count = 38
        table.Range.Row = 1
        table.Range.Rows.Count = 2
        width = _write_table(sheet, [{"service_group": "M2M", "No.": 1}], preserve_extra_columns=True)
        self.assertEqual(width, 38)
        self.assertEqual(len(ranges[((2,1),(2,len(MASTER_COLUMNS)))].Value[0]), len(MASTER_COLUMNS))
        table.Resize.assert_called_once_with(ranges[((1,1),(2,38))])

    def test_populated_extra_columns_need_explicit_mapping(self):
        sheet = MagicMock()
        sheet.ListObjects.return_value.ListColumns.Count = 38
        sheet.ListObjects.return_value.Range.Rows.Count = 2
        sheet.Range.return_value.Value2 = (("manual-data",None,None,None),)
        with self.assertRaisesRegex(ValueError, "populated extra columns"):
            _write_table(sheet, [], preserve_extra_columns=True)

    def test_extracts_four_weeks_full_precision_and_one_series_monthly_chart(self):
        def monthly(value):
            return {"Sep-26": (value,), "Oct-26": (value+1,)}, (value,)
        def weekly(value):
            return {f"W{week}": (value,) for week in range(37,41)}, (value,)
        workbook = MagicMock()
        workbook.__contains__.return_value = True
        with (
            patch("weekly.m2m_presentation.load_workbook", return_value=workbook),
            patch("weekly.m2m_presentation._section", return_value=1),
            patch("weekly.m2m_presentation._pivot_rows", side_effect=[
                monthly(0.06), monthly(40), monthly(61.2975),
                weekly(0.0601), weekly(5), weekly(2.5714), weekly(1.4),
            ]),
        ):
            charts, averages = extract_m2m_charts("m2m.xlsx", date(2026,10,1), 40)
        self.assertEqual(charts[1].labels, ("W37","W38","W39","W40"))
        self.assertEqual(charts[1].first, (2.5714,)*4)
        self.assertEqual(charts[2].second, (0.0601,)*4)
        self.assertEqual(charts[3].labels, ("SEP-26","OCT-26"))
        self.assertEqual(charts[3].first, (61.2975,62.2975))
        self.assertEqual(charts[3].second, ())
        self.assertEqual(averages["monthly_m2m_cash"], 61)
        workbook.close.assert_called_once()

    def test_single_series_cache_and_embedded_workbook_agree(self):
        namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        xml = f'''<c:chartSpace xmlns:c="{namespace}"><c:ser>
          <c:tx><c:v>M2M SERVICE</c:v></c:tx>
          <c:cat><c:strRef><c:f>Sheet1!B1:C1</c:f><c:strCache><c:ptCount val="2"/></c:strCache></c:strRef></c:cat>
          <c:val><c:numRef><c:f>Sheet1!B2:C2</c:f><c:numCache><c:ptCount val="2"/></c:numCache></c:numRef></c:val>
          <c:dLbls><c:dLblPos val="b"/></c:dLbls>
        </c:ser></c:chartSpace>'''.encode()
        data = ChartData(("SEP-26","OCT-26"), (1.4,61.2975), ())
        updated = etree.fromstring(_update_chart_xml(xml,data,2,3,
            expected_series=EXPECTED_SERIES[3],weekly=False,label_position="t"))
        self.assertEqual(updated.xpath("//c:val//c:pt/c:v/text()", namespaces={"c":namespace}), ["1.4","61.2975"])
        self.assertEqual(updated.xpath("//c:dLblPos/@val", namespaces={"c":namespace}), ["t"])
        source = BytesIO()
        workbook = Workbook()
        workbook.active.cell(3,2,999)
        workbook.save(source)
        result, first = _update_embedded(source.getvalue(),data,weekly=False,report_week=40,
                                        number_formats=("#,##0",))
        edited = load_workbook(BytesIO(result))
        self.addCleanup(edited.close)
        self.assertEqual(first, 2)
        self.assertEqual(edited.active["C2"].value,61.2975)
        self.assertIsNone(edited.active["B3"].value)

    def test_sparse_periods_use_zero_counts_and_gaps_for_undefined_means(self):
        workbook = MagicMock()
        workbook.__contains__.return_value = True
        months = ({"Sep-26": (1.5,)}, (1.5,))
        weeks = ({"W39": (1.5,)}, (1.5,))
        with (
            patch("weekly.m2m_presentation.load_workbook", return_value=workbook),
            patch("weekly.m2m_presentation._section", return_value=1),
            patch("weekly.m2m_presentation._pivot_rows", side_effect=[months]*3 + [weeks]*4),
        ):
            charts, _ = extract_m2m_charts("m2m.xlsx",date(2026,10,1),40)
        self.assertEqual(charts[1].first, (None,None,1.5,None))
        self.assertEqual(charts[2].first, (0,0,1.5,0))
        self.assertEqual(charts[3].first, (1.5,None))
        self.assertEqual(charts[4].first, (1.5,0))

    def test_blank_cache_points_remain_gaps_and_na_annotation_is_reusable(self):
        namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
        from weekly.presentation import _replace_cache
        cache = etree.fromstring(f'<c:numCache xmlns:c="{namespace}"><c:ptCount val="0"/></c:numCache>')
        _replace_cache(cache,(1.5,None,2))
        self.assertEqual(cache.xpath("c:pt/@idx",namespaces={"c":namespace}), ["0","2"])
        self.assertEqual(cache.xpath("c:ptCount/@val",namespaces={"c":namespace}), ["3"])
        shape = etree.fromstring('<root xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>M2M SERVICE :16</a:t></root>')
        with patch("weekly.presentation._shape", return_value=shape):
            _set_annotation(shape,"15",(None,))
            self.assertEqual(shape[0].text, "M2M SERVICE :N/A")
            _set_annotation(shape,"15",(2,))
            self.assertEqual(shape[0].text, "M2M SERVICE :2")


if __name__ == "__main__":
    unittest.main()
